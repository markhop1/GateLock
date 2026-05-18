#!/usr/bin/env python3
"""
Evaluación completa del sistema de reconocimiento facial.

Compara dos bases de datos (sin aug vs. con aug) usando probes de amigos reales
e impostores del dataset LFW. Barre thresholds para construir curvas ROC y exporta
todos los resultados a un archivo Excel con cuatro hojas.

Estructura de directorios esperada:
    evaluation/
        enrollment/{nombre}/  ← fotos de enrollment (para construir la DB)
        probes/{nombre}/      ← fotos de test (distintas de enrollment)
        impostors/{nombre}/   ← personas desconocidas (LFW)

Uso:
    python scripts/evaluate_recognition.py
    python scripts/evaluate_recognition.py --enrollment-dir mi_dir/enrollment ...
    python scripts/evaluate_recognition.py --skip-build   # reutiliza DBs existentes
"""

import argparse
import glob
import logging
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import cv2
import numpy as np

# ---------------------------------------------------------------------------
# Configurar path para importar módulos del proyecto
# ---------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(PROJECT_DIR))

from config import AUG_PER_IMAGE
from detection.face_detector import RetinaFaceDetector
from recognition.face_recognizer import MobileFaceNetRecognizer

logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")
logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Helpers de imagen
# ---------------------------------------------------------------------------

IMAGE_EXTS = ("*.jpg", "*.jpeg", "*.png", "*.webp", "*.bmp")


def list_images(directory: Path) -> List[Path]:
    files: List[Path] = []
    for ext in IMAGE_EXTS:
        files.extend(directory.glob(ext))
    return sorted(files)


def l2_normalize(x: np.ndarray, eps: float = 1e-12) -> np.ndarray:
    return x / (np.linalg.norm(x) + eps)


# ---------------------------------------------------------------------------
# Extracción de embeddings de una carpeta completa
# ---------------------------------------------------------------------------

def extract_embeddings_from_dir(
    root_dir: Path,
    detector: RetinaFaceDetector,
    fallback_recognizer: MobileFaceNetRecognizer,
) -> Dict[str, List[Tuple[Path, np.ndarray]]]:
    """
    Para cada sub-carpeta (identidad) en root_dir extrae un embedding por imagen.

    Returns:
        dict: {identity_name: [(image_path, embedding), ...]}
    """
    results: Dict[str, List[Tuple[Path, np.ndarray]]] = {}

    identity_dirs = sorted([d for d in root_dir.iterdir() if d.is_dir()])
    if not identity_dirs:
        logger.warning(f"No se encontraron sub-carpetas de identidad en {root_dir}")
        return results

    for id_dir in identity_dirs:
        name = id_dir.name
        images = list_images(id_dir)
        if not images:
            logger.warning(f"  {name}: sin imágenes, omitido")
            continue

        collected: List[Tuple[Path, np.ndarray]] = []
        for img_path in images:
            img_bgr = cv2.imread(str(img_path))
            if img_bgr is None:
                continue

            try:
                detections = detector.detect_faces(img_bgr)
            except Exception as e:
                logger.debug(f"Detección fallida para {img_path}: {e}")
                continue

            if not detections:
                logger.debug(f"Sin detección en {img_path}")
                continue

            largest = max(detections, key=lambda d: d.width * d.height)

            if getattr(largest, "embedding", None) is not None:
                emb = l2_normalize(np.asarray(largest.embedding).flatten())
            else:
                face_crop = largest.get_face_crop(img_bgr)
                if face_crop.size == 0:
                    continue
                face_rgb = cv2.cvtColor(face_crop, cv2.COLOR_BGR2RGB)
                emb = l2_normalize(fallback_recognizer.extract_embedding(face_rgb))

            collected.append((img_path, emb))

        logger.info(f"  {name}: {len(collected)}/{len(images)} imágenes con embedding")
        if collected:
            results[name] = collected

    return results


# ---------------------------------------------------------------------------
# Construir DB (mean embedding por identidad) y guardar .npy
# ---------------------------------------------------------------------------

def build_db(
    enrollment_dir: Path,
    output_dir: Path,
    detector: RetinaFaceDetector,
    fallback_recognizer: MobileFaceNetRecognizer,
    use_augment: bool,
    aug_per_image: int,
) -> Tuple[np.ndarray, np.ndarray]:
    """
    Construye una base de datos de embeddings y la guarda en output_dir.

    Returns:
        (embeddings_array [N, 512], labels_array [N])
    """
    import albumentations as A
    from insightface.utils import face_align

    def build_augmenter() -> A.Compose:
        return A.Compose([
            A.RandomBrightnessContrast(brightness_limit=0.35, contrast_limit=0.25, p=0.9),
            A.RandomGamma(gamma_limit=(70, 150), p=0.7),
            A.HueSaturationValue(hue_shift_limit=5, sat_shift_limit=20, val_shift_limit=10, p=0.6),
            A.OneOf([
                A.GaussianBlur(blur_limit=(3, 7), p=1.0),
                A.MotionBlur(blur_limit=7, p=1.0),
            ], p=0.25),
            A.GaussNoise(std_range=(0.01, 0.05), mean_range=(0.0, 0.0), per_channel=True, p=0.25),
            A.ImageCompression(quality_range=(35, 95), compression_type="jpeg", p=0.25),
        ])

    def quality_gate(face_rgb: np.ndarray) -> bool:
        gray = cv2.cvtColor(face_rgb, cv2.COLOR_RGB2GRAY)
        mean = float(gray.mean())
        if mean < 35 or mean > 220:
            return False
        if float((gray < 10).mean()) > 0.25:
            return False
        if float((gray > 245).mean()) > 0.25:
            return False
        return True

    augmenter = build_augmenter() if use_augment else None
    rec_model = detector.app.models.get("recognition") if use_augment else None

    identity_dirs = sorted([d for d in enrollment_dir.iterdir() if d.is_dir()])
    per_id_agg: Dict[str, np.ndarray] = {}

    for id_dir in identity_dirs:
        name = id_dir.name
        images = list_images(id_dir)
        if not images:
            continue

        collected: List[np.ndarray] = []
        for img_path in images:
            img_bgr = cv2.imread(str(img_path))
            if img_bgr is None:
                continue

            try:
                detections = detector.detect_faces(img_bgr)
            except Exception:
                continue

            if not detections:
                continue

            largest = max(detections, key=lambda d: d.width * d.height)

            if getattr(largest, "embedding", None) is not None:
                emb = l2_normalize(np.asarray(largest.embedding).flatten())
            else:
                face_crop = largest.get_face_crop(img_bgr)
                if face_crop.size == 0:
                    continue
                face_rgb_fallback = cv2.cvtColor(face_crop, cv2.COLOR_BGR2RGB)
                emb = l2_normalize(fallback_recognizer.extract_embedding(face_rgb_fallback))

            collected.append(emb)

            # Augmentation sobre normed crop — solo con --augment
            if (
                augmenter is not None
                and rec_model is not None
                and getattr(largest, "kps", None) is not None
            ):
                try:
                    img_rgb = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2RGB)
                    aimg_rgb = face_align.norm_crop(img_rgb, landmark=largest.kps, image_size=112)
                    for _ in range(aug_per_image):
                        aug_rgb = augmenter(image=aimg_rgb)["image"]
                        if not quality_gate(aug_rgb):
                            continue
                        aug_bgr = cv2.cvtColor(aug_rgb, cv2.COLOR_RGB2BGR)
                        aug_emb = rec_model.get_feat([aug_bgr])[0]
                        collected.append(l2_normalize(aug_emb.flatten()))
                except Exception as aug_err:
                    logger.debug(f"Augmentación fallida para {img_path}: {aug_err}")

        if collected:
            mat = np.stack(collected, axis=0)
            per_id_agg[name] = l2_normalize(mat.mean(axis=0))
            logger.info(
                f"  [{name}] {len(collected)} embeddings → 1 embedding agregado"
                + (" (incl. aug)" if use_augment else "")
            )

    if not per_id_agg:
        raise RuntimeError(f"No se generaron embeddings desde {enrollment_dir}")

    id_list = sorted(per_id_agg.keys())
    embeddings_array = np.stack([per_id_agg[i] for i in id_list], axis=0)
    labels_array = np.array(id_list)

    suffix = "_aug" if use_augment else "_noaug"
    output_dir.mkdir(parents=True, exist_ok=True)
    np.save(output_dir / f"face_embeddings{suffix}.npy", embeddings_array)
    np.save(output_dir / f"face_labels{suffix}.npy", labels_array)
    logger.info(f"  DB guardada en {output_dir} (sufijo: {suffix})")

    return embeddings_array, labels_array


# ---------------------------------------------------------------------------
# Evaluación: sweep de thresholds
# ---------------------------------------------------------------------------

def compute_similarity(query_emb: np.ndarray, db_embeddings: np.ndarray) -> Tuple[int, float]:
    """Devuelve (best_idx, best_cosine_similarity)."""
    query_emb = l2_normalize(query_emb.flatten())
    sims = db_embeddings @ query_emb
    best_idx = int(np.argmax(sims))
    return best_idx, float(sims[best_idx])


RawRow = dict  # typing alias


def run_probes(
    probe_embs: Dict[str, List[Tuple[Path, np.ndarray]]],
    impostor_embs: Dict[str, List[Tuple[Path, np.ndarray]]],
    db_embeddings: np.ndarray,
    db_labels: np.ndarray,
    mode_label: str,
    enrolled_names: set,
) -> List[RawRow]:
    """
    Genera filas de resultados crudos (una por probe image × mode).

    true_label = 1  si la identidad del probe está en la DB (amigo)
    true_label = 0  si es impostor
    """
    rows: List[RawRow] = []

    def process(name: str, img_path: Path, emb: np.ndarray, true_label: int) -> RawRow:
        best_idx, score = compute_similarity(emb, db_embeddings)
        predicted_name = str(db_labels[best_idx])
        return {
            "probe_id": img_path.name,
            "probe_path": str(img_path),
            "true_name": name,
            "true_label": true_label,
            "predicted_name": predicted_name,
            "score": round(score, 6),
            "augmentation_mode": mode_label,
        }

    for name, items in probe_embs.items():
        tl = 1 if name in enrolled_names else 0
        for img_path, emb in items:
            rows.append(process(name, img_path, emb, tl))

    for name, items in impostor_embs.items():
        for img_path, emb in items:
            rows.append(process(name, img_path, emb, 0))

    return rows


def compute_roc_data(
    rows: List[RawRow],
    thresholds: np.ndarray,
) -> List[dict]:
    """
    Para cada threshold calcula FPR, TPR, FAR (=FPR), FRR (=1-TPR).
    Un probe de true_label=1 es positivo; de true_label=0 es negativo.
    Match = score >= threshold.
    """
    scores = np.array([r["score"] for r in rows])
    true_labels = np.array([r["true_label"] for r in rows])

    roc_rows = []
    for thr in thresholds:
        predicted = (scores >= thr).astype(int)

        tp = int(((predicted == 1) & (true_labels == 1)).sum())
        fp = int(((predicted == 1) & (true_labels == 0)).sum())
        tn = int(((predicted == 0) & (true_labels == 0)).sum())
        fn = int(((predicted == 0) & (true_labels == 1)).sum())

        tpr = tp / (tp + fn) if (tp + fn) > 0 else 0.0
        fpr = fp / (fp + tn) if (fp + tn) > 0 else 0.0
        roc_rows.append({
            "threshold": round(float(thr), 4),
            "TPR": round(tpr, 6),
            "FPR": round(fpr, 6),
            "FAR": round(fpr, 6),
            "FRR": round(1.0 - tpr, 6),
            "TP": tp, "FP": fp, "TN": tn, "FN": fn,
        })
    return roc_rows


def compute_per_person_metrics(
    rows: List[RawRow],
    threshold: float,
    enrolled_names: set,
) -> List[dict]:
    """Métricas por persona enrolled al threshold dado."""
    person_metrics = []
    for name in sorted(enrolled_names):
        person_rows = [r for r in rows if r["true_name"] == name]
        if not person_rows:
            continue
        scores = np.array([r["score"] for r in person_rows])
        predicted = (scores >= threshold).astype(int)
        true_labels = np.array([r["true_label"] for r in person_rows])

        tp = int(((predicted == 1) & (true_labels == 1)).sum())
        fp = int(((predicted == 1) & (true_labels == 0)).sum())
        tn = int(((predicted == 0) & (true_labels == 0)).sum())
        fn = int(((predicted == 0) & (true_labels == 1)).sum())

        precision = tp / (tp + fp) if (tp + fp) > 0 else 0.0
        recall = tp / (tp + fn) if (tp + fn) > 0 else 0.0
        f1 = (2 * precision * recall / (precision + recall)) if (precision + recall) > 0 else 0.0
        far = fp / (fp + tn) if (fp + tn) > 0 else 0.0
        frr = fn / (fn + tp) if (fn + tp) > 0 else 0.0

        person_metrics.append({
            "person": name,
            "threshold": round(threshold, 4),
            "TP": tp, "FP": fp, "TN": tn, "FN": fn,
            "FAR": round(far, 6),
            "FRR": round(frr, 6),
            "precision": round(precision, 6),
            "recall": round(recall, 6),
            "F1": round(f1, 6),
        })
    return person_metrics


def compute_auc(roc_rows: List[dict]) -> float:
    """AUC por método trapezoidal."""
    sorted_rows = sorted(roc_rows, key=lambda r: r["FPR"])
    fprs = np.array([r["FPR"] for r in sorted_rows])
    tprs = np.array([r["TPR"] for r in sorted_rows])
    return float(np.trapz(tprs, fprs))


def find_eer(roc_rows: List[dict]) -> Tuple[float, float]:
    """Equal Error Rate: punto donde FAR ≈ FRR. Devuelve (eer, threshold)."""
    diffs = [abs(r["FAR"] - r["FRR"]) for r in roc_rows]
    idx = int(np.argmin(diffs))
    row = roc_rows[idx]
    eer = (row["FAR"] + row["FRR"]) / 2.0
    return round(eer, 6), row["threshold"]


def find_best_threshold(roc_rows: List[dict]) -> Tuple[float, float]:
    """Threshold con mayor (TPR - FPR), i.e. máximo Youden J."""
    best = max(roc_rows, key=lambda r: r["TPR"] - r["FPR"])
    return round(best["TPR"] - best["FPR"], 6), best["threshold"]


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def export_to_excel(
    output_path: Path,
    raw_noaug: List[RawRow],
    raw_aug: List[RawRow],
    roc_noaug: List[dict],
    roc_aug: List[dict],
    per_person_noaug: List[dict],
    per_person_aug: List[dict],
    summary: List[dict],
) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("openpyxl no está instalado. Instala con: pip install openpyxl")
        raise

    wb = openpyxl.Workbook()

    HEADER_FONT = Font(bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(fill_type="solid", fgColor="2F5496")
    AUG_FILL = PatternFill(fill_type="solid", fgColor="E2EFDA")    # verde claro
    NOAUG_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")  # azul claro

    def style_header_row(ws, col_count: int):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

    def auto_width(ws):
        for col in ws.columns:
            max_len = max((len(str(c.value)) for c in col if c.value), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 40)

    # ── Sheet 1: Raw Results ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Raw Results"
    if raw_noaug or raw_aug:
        all_raw = raw_noaug + raw_aug
        headers = list(all_raw[0].keys())
        ws1.append(headers)
        style_header_row(ws1, len(headers))
        for row in all_raw:
            ws1.append(list(row.values()))
    auto_width(ws1)

    # ── Sheet 2: Per Person Metrics ───────────────────────────────────────
    ws2 = wb.create_sheet("Per Person Metrics")
    all_person = (
        [{"mode": "no_aug", **r} for r in per_person_noaug]
        + [{"mode": "aug", **r} for r in per_person_aug]
    )
    if all_person:
        headers2 = list(all_person[0].keys())
        ws2.append(headers2)
        style_header_row(ws2, len(headers2))
        for i, row in enumerate(all_person, start=2):
            ws2.append(list(row.values()))
            fill = AUG_FILL if row["mode"] == "aug" else NOAUG_FILL
            for col in range(1, len(headers2) + 1):
                ws2.cell(row=i, column=col).fill = fill
    auto_width(ws2)

    # ── Sheet 3: ROC Curve Data ───────────────────────────────────────────
    ws3 = wb.create_sheet("ROC Curve Data")
    roc_headers = (
        ["threshold_noaug", "TPR_noaug", "FPR_noaug", "FAR_noaug", "FRR_noaug",
         "TP_noaug", "FP_noaug", "TN_noaug", "FN_noaug",
         "threshold_aug", "TPR_aug", "FPR_aug", "FAR_aug", "FRR_aug",
         "TP_aug", "FP_aug", "TN_aug", "FN_aug"]
    )
    ws3.append(roc_headers)
    style_header_row(ws3, len(roc_headers))
    for na, au in zip(roc_noaug, roc_aug):
        ws3.append([
            na["threshold"], na["TPR"], na["FPR"], na["FAR"], na["FRR"],
            na["TP"], na["FP"], na["TN"], na["FN"],
            au["threshold"], au["TPR"], au["FPR"], au["FAR"], au["FRR"],
            au["TP"], au["FP"], au["TN"], au["FN"],
        ])
    auto_width(ws3)

    # ── Sheet 4: Summary Comparison ───────────────────────────────────────
    ws4 = wb.create_sheet("Summary Comparison")
    if summary:
        s_headers = list(summary[0].keys())
        ws4.append(s_headers)
        style_header_row(ws4, len(s_headers))
        for i, row in enumerate(summary, start=2):
            ws4.append(list(row.values()))
            fill = AUG_FILL if row.get("mode") == "aug" else NOAUG_FILL
            for col in range(1, len(s_headers) + 1):
                ws4.cell(row=i, column=col).fill = fill
    auto_width(ws4)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    logger.info(f"Excel exportado a: {output_path}")


# ---------------------------------------------------------------------------
# Entrada principal
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    eval_dir = PROJECT_DIR / "evaluation"
    parser = argparse.ArgumentParser(
        description="Evalúa el sistema de reconocimiento facial y exporta resultados a Excel.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--enrollment-dir", type=Path, default=eval_dir / "enrollment",
                        help="Carpeta con sub-carpetas por persona para construir la DB")
    parser.add_argument("--probes-dir", type=Path, default=eval_dir / "probes",
                        help="Carpeta con sub-carpetas por persona para test")
    parser.add_argument("--impostors-dir", type=Path, default=eval_dir / "impostors",
                        help="Carpeta con sub-carpetas de impostores (LFW)")
    parser.add_argument("--db-dir", type=Path, default=eval_dir / "db",
                        help="Carpeta donde se guardan/leen las DBs (.npy)")
    parser.add_argument("--output-excel", type=Path, default=eval_dir / "results.xlsx",
                        help="Ruta del Excel de salida")
    parser.add_argument("--aug-per-image", type=int, default=AUG_PER_IMAGE,
                        help="Versiones augmentadas por imagen al construir DB con aug")
    parser.add_argument("--threshold-steps", type=int, default=200,
                        help="Número de thresholds en el sweep ROC")
    parser.add_argument("--default-threshold", type=float, default=0.35,
                        help="Threshold fijo para métricas por persona")
    parser.add_argument("--skip-build", action="store_true",
                        help="Reutilizar archivos .npy ya existentes en --db-dir")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    # Validar directorios
    for attr, label in [
        ("enrollment_dir", "enrollment"),
        ("probes_dir", "probes"),
        ("impostors_dir", "impostors"),
    ]:
        d = getattr(args, attr)
        if not d.exists():
            logger.error(f"Directorio {label} no encontrado: {d}")
            sys.exit(1)

    # Inicializar modelos
    logger.info("Inicializando modelos InsightFace...")
    detector = RetinaFaceDetector()
    detector.initialize()
    fallback_rec = MobileFaceNetRecognizer()
    fallback_rec.initialize()

    # Construir o cargar DBs
    emb_noaug_path = args.db_dir / "face_embeddings_noaug.npy"
    lbl_noaug_path = args.db_dir / "face_labels_noaug.npy"
    emb_aug_path = args.db_dir / "face_embeddings_aug.npy"
    lbl_aug_path = args.db_dir / "face_labels_aug.npy"

    if args.skip_build and emb_noaug_path.exists() and emb_aug_path.exists():
        logger.info("--skip-build: cargando DBs existentes...")
        db_emb_noaug = np.load(emb_noaug_path)
        db_lbl_noaug = np.load(lbl_noaug_path)
        db_emb_aug = np.load(emb_aug_path)
        db_lbl_aug = np.load(lbl_aug_path)
    else:
        logger.info("Construyendo DB sin augmentación...")
        db_emb_noaug, db_lbl_noaug = build_db(
            args.enrollment_dir, args.db_dir, detector, fallback_rec,
            use_augment=False, aug_per_image=args.aug_per_image,
        )
        logger.info("Construyendo DB con augmentación...")
        db_emb_aug, db_lbl_aug = build_db(
            args.enrollment_dir, args.db_dir, detector, fallback_rec,
            use_augment=True, aug_per_image=args.aug_per_image,
        )

    enrolled_noaug = set(db_lbl_noaug.tolist())
    enrolled_aug = set(db_lbl_aug.tolist())

    # Extraer embeddings de probes e impostores (una sola vez, independiente de la DB)
    logger.info("Extrayendo embeddings de probes...")
    probe_embs = extract_embeddings_from_dir(args.probes_dir, detector, fallback_rec)

    logger.info("Extrayendo embeddings de impostores...")
    impostor_embs = extract_embeddings_from_dir(args.impostors_dir, detector, fallback_rec)

    if not probe_embs and not impostor_embs:
        logger.error("No se encontraron embeddings en probes ni impostores. Abortando.")
        sys.exit(1)

    # Generar filas crudas
    thresholds = np.linspace(0.0, 1.0, args.threshold_steps)

    logger.info("Evaluando con DB sin augmentación...")
    raw_noaug = run_probes(probe_embs, impostor_embs, db_emb_noaug, db_lbl_noaug, "no_aug", enrolled_noaug)
    logger.info("Evaluando con DB con augmentación...")
    raw_aug = run_probes(probe_embs, impostor_embs, db_emb_aug, db_lbl_aug, "aug", enrolled_aug)

    # ROC
    roc_noaug = compute_roc_data(raw_noaug, thresholds)
    roc_aug = compute_roc_data(raw_aug, thresholds)

    # Per person
    per_person_noaug = compute_per_person_metrics(raw_noaug, args.default_threshold, enrolled_noaug)
    per_person_aug = compute_per_person_metrics(raw_aug, args.default_threshold, enrolled_aug)

    # Summary
    auc_noaug = compute_auc(roc_noaug)
    auc_aug = compute_auc(roc_aug)
    eer_noaug, eer_thr_noaug = find_eer(roc_noaug)
    eer_aug, eer_thr_aug = find_eer(roc_aug)
    youden_noaug, best_thr_noaug = find_best_threshold(roc_noaug)
    youden_aug, best_thr_aug = find_best_threshold(roc_aug)

    summary = [
        {
            "mode": "no_aug",
            "AUC": round(auc_noaug, 6),
            "EER": eer_noaug,
            "EER_threshold": eer_thr_noaug,
            "best_youden_J": youden_noaug,
            "best_threshold_youden": best_thr_noaug,
            "default_threshold": args.default_threshold,
            "total_probes": len(raw_noaug),
            "enrolled_people": len(enrolled_noaug),
        },
        {
            "mode": "aug",
            "AUC": round(auc_aug, 6),
            "EER": eer_aug,
            "EER_threshold": eer_thr_aug,
            "best_youden_J": youden_aug,
            "best_threshold_youden": best_thr_aug,
            "default_threshold": args.default_threshold,
            "total_probes": len(raw_aug),
            "enrolled_people": len(enrolled_aug),
        },
    ]

    logger.info(f"\n=== Resumen ===")
    logger.info(f"  AUC  no_aug={auc_noaug:.4f}  aug={auc_aug:.4f}")
    logger.info(f"  EER  no_aug={eer_noaug:.4f} @ thr={eer_thr_noaug}  |  aug={eer_aug:.4f} @ thr={eer_thr_aug}")
    logger.info(f"  Best threshold (Youden)  no_aug={best_thr_noaug}  aug={best_thr_aug}")

    # Exportar Excel
    export_to_excel(
        output_path=args.output_excel,
        raw_noaug=raw_noaug,
        raw_aug=raw_aug,
        roc_noaug=roc_noaug,
        roc_aug=roc_aug,
        per_person_noaug=per_person_noaug,
        per_person_aug=per_person_aug,
        summary=summary,
    )


if __name__ == "__main__":
    main()
